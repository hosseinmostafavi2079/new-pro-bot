import os
from datetime import datetime
from openpyxl import Workbook, load_workbook

# --- این متغیر حتما باید اینجا تعریف شده باشد ---
LOG_FILE = "products_log.xlsx"
# ---------------------------------------------

def _get_sheet():
    """تابع کمکی برای باز کردن یا ساختن فایل اکسل"""
    if not os.path.exists(LOG_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Logs"
        # هدرهای اصلی
        ws.append(["تاریخ و ساعت", "آیدی ادمین", "نام محصول", "عملیات / قیمت", "SKU", "جزئیات / لینک"])
        return wb, ws
    else:
        wb = load_workbook(LOG_FILE)
        ws = wb.active
        return wb, ws

def log_product_to_excel(user_id, title, price, sku, url):
    """ثبت محصول جدید"""
    try:
        wb, ws = _get_sheet()
        current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # افزودن سطر: عملیات = قیمت، جزئیات = لینک
        ws.append([current_time, user_id, title, f"{price} تومان", sku, url])
        
        wb.save(LOG_FILE)
        print(f"LOG: Added product {sku}")
        return True
    except Exception as e:
        print(f"⚠️ Excel Log Error: {e}")
        return False

def log_edit_to_excel(user_id, title, sku, action_type, details):
    """ثبت ویرایش محصول"""
    try:
        wb, ws = _get_sheet()
        current_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # در ستون قیمت، نوع عملیات را می‌نویسیم
        # در ستون لینک، جزئیات را می‌نویسیم
        ws.append([current_time, user_id, title, action_type, sku, details])
        
        wb.save(LOG_FILE)
        print(f"LOG: Edited product {sku}")
        return True
    except Exception as e:
        print(f"⚠️ Excel Edit Log Error: {e}")
        return False